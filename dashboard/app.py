import os
from functools import wraps
from flask import Flask, render_template, jsonify, request, session
from werkzeug.security import check_password_hash, generate_password_hash

from database import get_connection, init_db, db_exists, migrate_db, REGIONAL_ORDER

app = Flask(__name__)
app.secret_key = os.environ.get('SECRET_KEY', 'bantek-dashboard-secret-key-change-me')


@app.errorhandler(404)
def not_found(e):
    return jsonify({'error': 'Not Found', 'message': str(e)}), 404


@app.errorhandler(500)
def server_error(e):
    return jsonify({'error': 'Server Error', 'message': 'Terjadi kesalahan pada server.'}), 500

KEGIATAN_MAP = [
    ('Tanam Ulang', 'Tanam Ulang / Program TU'),
    ('Program TU', 'Tanam Ulang / Program TU'),
    ('Bibit', 'Bibit / Lewat Umur'),
    ('Lewat Umur', 'Bibit / Lewat Umur'),
    ('Non Produktif', 'Areal Non Produktif / Puso'),
    ('Puso', 'Areal Non Produktif / Puso'),
    ('Produksi', 'Produksi / Kultur Teknis'),
    ('Kultur Teknis', 'Produksi / Kultur Teknis'),
    ('Kajian Kelayakan', 'BT / FS (Kajian Kelayakan)'),
    ('Evaluasi', 'Evaluasi/Kajian/Pemeriksaan'),
    ('Pemeriksaan', 'Evaluasi/Kajian/Pemeriksaan'),
]


def classify_kegiatan(kegiatan):
    k = str(kegiatan) if kegiatan else ''
    for keyword, category in KEGIATAN_MAP:
        if keyword.lower() in k.lower():
            return category
    return 'Lainnya' if k else '(Tidak ada kegiatan)'


def status_of(sudah_dikirim):
    return 'SELESAI' if sudah_dikirim and str(sudah_dikirim).strip() != '' else 'PROSES'


def ensure_db():
    if not db_exists():
        raise RuntimeError("Database belum ada. Jalankan: python init_db.py")


def fetch_laporan(conn, id):
    row = conn.execute("SELECT * FROM laporan WHERE id=?", (id,)).fetchone()
    return dict(row) if row else None


def laporan_public(rec):
    rec = dict(rec)
    rec['status'] = status_of(rec.get('sudah_dikirim'))
    rec['kegiatan_kategori'] = classify_kegiatan(rec.get('kegiatan'))
    return rec


def query_all_laporan():
    conn = get_connection()
    try:
        rows = conn.execute("SELECT * FROM laporan").fetchall()
        return [laporan_public(r) for r in rows]
    finally:
        conn.close()


def _regional_no(regional):
    m = {'Regional 1': '1', 'Regional 2': '2', 'Regional 3': '3',
         'Regional 4': '4', 'Regional 5': '5', 'Regional 6': '6',
         'Regional 7': '7', 'Regional 4 Palmco': '4P', 'PT Swasta / PPKS': 'SW'}
    return m.get(regional, '0')


def make_kode(rec, seq):
    no_reg = _regional_no(rec.get('regional'))
    tahun = str(rec.get('tahun') or '')
    if not tahun:
        tahun = '0000'
    return 'R%s-%s-%03d' % (no_reg, tahun, seq)


def _backfill_single(id):
    conn = get_connection()
    try:
        rec = fetch_laporan(conn, id)
        if not rec or rec.get('kode_laporan'):
            return
        no_reg = _regional_no(rec.get('regional'))
        tahun = str(rec.get('tahun') or '') or '0000'
        row = conn.execute(
            "SELECT COUNT(*) AS c FROM laporan WHERE regional=? AND kode_laporan IS NOT NULL AND kode_laporan != ''",
            (rec.get('regional'),)).fetchone()
        seq = row['c'] + 1
        kode = 'R%s-%s-%03d' % (no_reg, tahun, seq)
        conn.execute("UPDATE laporan SET kode_laporan=? WHERE id=?", (kode, id))
        conn.commit()
        return kode
    finally:
        conn.close()


def migrate_and_backfill():
    conn = get_connection()
    try:
        cols = [r['name'] for r in conn.execute("PRAGMA table_info(laporan)").fetchall()]
        if 'kode_laporan' not in cols:
            conn.execute("ALTER TABLE laporan ADD COLUMN kode_laporan TEXT")
            conn.commit()
        rows = conn.execute(
            "SELECT * FROM laporan WHERE kode_laporan IS NULL OR kode_laporan=''"
        ).fetchall()
        counters = {}
        seqs = {}
        for r in rows:
            d = dict(r)
            reg = d['regional']
            counters[reg] = counters.get(reg, 0) + 1
            no_reg = _regional_no(reg)
            tahun = str(d.get('tahun') or '') or '0000'
            key = (reg, tahun)
            seqs[key] = seqs.get(key, 0) + 1
            kode = 'R%s-%s-%03d' % (no_reg, tahun, seqs[key])
            conn.execute("UPDATE laporan SET kode_laporan=? WHERE id=?", (kode, d['id']))
        conn.commit()
        return len(rows)
    finally:
        conn.close()


# ---------------- AUTH ----------------

def login_required(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        if not session.get('admin_logged_in'):
            return jsonify({'error': 'Unauthorized', 'message': 'Anda harus login terlebih dahulu.'}), 401
        return f(*args, **kwargs)
    return decorated


def check_admin_password(password):
    conn = get_connection()
    try:
        row = conn.execute("SELECT value FROM settings WHERE key='admin_password_hash'").fetchone()
        if not row:
            return False
        return check_password_hash(row['value'], password)
    finally:
        conn.close()


# ---------------- PAGES ----------------

@app.route('/')
def index():
    return render_template('index.html')


# ---------------- AUTH API ----------------

@app.route('/api/login', methods=['POST'])
def api_login():
    data = request.get_json(silent=True) or {}
    password = data.get('password', '') or request.form.get('password', '')
    if check_admin_password(password):
        session['admin_logged_in'] = True
        return jsonify({'status': 'ok'})
    return jsonify({'error': 'Password salah'}), 401


@app.route('/api/logout', methods=['POST'])
@login_required
def api_logout():
    session.pop('admin_logged_in', None)
    return jsonify({'status': 'ok'})


@app.route('/api/status')
def api_status():
    return jsonify({'logged_in': bool(session.get('admin_logged_in'))})


# ---------------- READ APIs (public) ----------------

@app.route('/api/summary')
def api_summary():
    data = query_all_laporan()
    total = len(data)
    selesai = sum(1 for d in data if d['status'] == 'SELESAI')
    proses = sum(1 for d in data if d['status'] == 'PROSES')
    persen = round(selesai / total * 100, 1) if total > 0 else 0
    return jsonify({'total': total, 'selesai': selesai, 'proses': proses, 'persen_selesai': persen})


@app.route('/api/regional')
def api_regional():
    data = query_all_laporan()
    regional_stats = {}
    for d in data:
        r = d['regional']
        if r not in regional_stats:
            regional_stats[r] = {'total': 0, 'selesai': 0, 'proses': 0}
        regional_stats[r]['total'] += 1
        if d['status'] == 'SELESAI':
            regional_stats[r]['selesai'] += 1
        else:
            regional_stats[r]['proses'] += 1
    result = []
    for r in REGIONAL_ORDER:
        if r in regional_stats:
            s = regional_stats[r]
            result.append({
                'regional': r, 'total': s['total'], 'selesai': s['selesai'],
                'proses': s['proses'],
                'persen': round(s['selesai'] / s['total'] * 100, 1) if s['total'] > 0 else 0,
            })
    return jsonify(result)


@app.route('/api/tahun')
def api_tahun():
    data = query_all_laporan()
    tahun_stats = {}
    for d in data:
        t = d['tahun']
        if not t:
            continue
        if t not in tahun_stats:
            tahun_stats[t] = {'total': 0, 'selesai': 0, 'proses': 0}
        tahun_stats[t]['total'] += 1
        if d['status'] == 'SELESAI':
            tahun_stats[t]['selesai'] += 1
        else:
            tahun_stats[t]['proses'] += 1
    result = []
    for t in sorted(tahun_stats.keys()):
        s = tahun_stats[t]
        result.append({
            'tahun': t, 'total': s['total'], 'selesai': s['selesai'],
            'proses': s['proses'],
            'persen': round(s['selesai'] / s['total'] * 100, 1) if s['total'] > 0 else 0,
        })
    return jsonify(result)


@app.route('/api/matriks_status')
def api_matriks_status():
    data = query_all_laporan()
    selesai_matrix = {r: {'2024': 0, '2025': 0, '2026': 0, 'total': 0} for r in REGIONAL_ORDER}
    proses_matrix = {r: {'2024': 0, '2025': 0, '2026': 0, 'total': 0} for r in REGIONAL_ORDER}
    for d in data:
        r, t = d['regional'], d['tahun']
        if r not in selesai_matrix or t not in ('2024', '2025', '2026'):
            continue
        if d['status'] == 'SELESAI':
            selesai_matrix[r][t] += 1
            selesai_matrix[r]['total'] += 1
        else:
            proses_matrix[r][t] += 1
            proses_matrix[r]['total'] += 1
    result = []
    for r in REGIONAL_ORDER:
        row = {'regional': r}
        for yr in ('2024', '2025', '2026', 'total'):
            t = selesai_matrix[r][yr] + proses_matrix[r][yr]
            s = selesai_matrix[r][yr]
            row[yr] = s
            row[yr + '_total'] = t
            row[yr + '_persen'] = round(s / t * 100) if t > 0 else 0
        result.append(row)
    return jsonify(result)


@app.route('/api/kegiatan')
def api_kegiatan():
    data = query_all_laporan()
    kegiatan_count = {}
    for d in data:
        k = d['kegiatan_kategori']
        kegiatan_count[k] = kegiatan_count.get(k, 0) + 1
    result = [{'kegiatan': k, 'count': v} for k, v in kegiatan_count.items()]
    result.sort(key=lambda x: x['count'], reverse=True)
    return jsonify(result)


@app.route('/api/korektor')
def api_korektor():
    data = query_all_laporan()
    korektor_stats = {}
    for d in data:
        korektor = d['korektor']
        if not korektor:
            continue
        names = [n.strip() for n in korektor.split(';') if n.strip()]
        for name in names:
            if name not in korektor_stats:
                korektor_stats[name] = {'total': 0, 'selesai': 0, 'proses': 0, 'regional': {r: 0 for r in REGIONAL_ORDER}}
            korektor_stats[name]['total'] += 1
            korektor_stats[name]['regional'][d['regional']] += 1
            if d['status'] == 'SELESAI':
                korektor_stats[name]['selesai'] += 1
            else:
                korektor_stats[name]['proses'] += 1
    result = []
    for name, stats in sorted(korektor_stats.items(), key=lambda x: x[1]['total'], reverse=True):
        result.append({
            'korektor': name, 'total': stats['total'], 'selesai': stats['selesai'],
            'proses': stats['proses'],
            'persen': round(stats['selesai'] / stats['total'] * 100, 1) if stats['total'] > 0 else 0,
            'regional': stats['regional'],
        })
    return jsonify(result)


@app.route('/api/detail')
def api_detail():
    data = query_all_laporan()
    regional = request.args.get('regional', '')
    tahun = request.args.get('tahun', '')
    status = request.args.get('status', '')
    search = request.args.get('search', '').lower()
    page = int(request.args.get('page', 1))
    per_page = int(request.args.get('per_page', 20))

    filtered = data
    if regional and regional != 'Semua':
        filtered = [d for d in filtered if d['regional'] == regional]
    if tahun and tahun != 'Semua':
        filtered = [d for d in filtered if d['tahun'] == tahun]
    if status and status != 'Semua':
        filtered = [d for d in filtered if d['status'] == status]
    if search:
        filtered = [d for d in filtered if
                    search in str(d['kebun']).lower() or
                    search in str(d['kegiatan']).lower() or
                    search in str(d['korektor']).lower() or
                    search in str(d['perusahaan']).lower() or
                    search in str(d['petugas']).lower() or
                    search in str(d['kode_laporan']).lower() or
                    search in str(d['regional']).lower() or
                    search in str(d['tahun']).lower()]

    total_filtered = len(filtered)
    start = (page - 1) * per_page
    end = start + per_page
    page_data = filtered[start:end]

    return jsonify({
        'total': total_filtered, 'page': page, 'per_page': per_page,
        'total_pages': (total_filtered + per_page - 1) // per_page,
        'data': page_data,
    })


# ---------------- WRITE APIs (login required) ----------------

def clean_str(v):
    return str(v).strip() if v is not None else ''


def parse_payload(payload, existing=None):
    existing = existing or {}
    def field(key, default=''):
        v = payload.get(key)
        if v is None:
            v = existing.get(key, default)
        return clean_str(v)
    data = {
        'kode_laporan': field('kode_laporan'),
        'no': field('no'),
        'regional': clean_str(payload.get('regional')) or clean_str(existing.get('regional')) or 'Regional 1',
        'perusahaan': field('perusahaan'),
        'kebun': field('kebun'),
        'petugas': field('petugas'),
        'tanggal_kunjungan': field('tanggal_kunjungan'),
        'kegiatan': field('kegiatan'),
        'draft_masuk': field('draft_masuk'),
        'draft_korektor': field('draft_korektor'),
        'korektor': field('korektor'),
        'revisi': field('revisi'),
        'cetak': field('cetak'),
        'sudah_dikirim': field('sudah_dikirim'),
        'tahun': field('tahun'),
        'folder_laporan': field('folder_laporan'),
        'catatan': field('catatan'),
    }
    status = clean_str(payload.get('status'))
    if status:
        if status.upper() == 'SELESAI' and not data['sudah_dikirim']:
            data['sudah_dikirim'] = '✓'
        elif status.upper() == 'PROSES':
            data['sudah_dikirim'] = ''
    return data


def insert_korektor_master(name):
    if name:
        conn = get_connection()
        try:
            conn.execute("INSERT OR IGNORE INTO korektor_master (nama) VALUES (?)", (name,))
            conn.commit()
        finally:
            conn.close()


def lookup_record(id):
    conn = get_connection()
    try:
        return fetch_laporan(conn, id)
    finally:
        conn.close()


@app.route('/api/laporan/<int:id>', methods=['GET'])
def api_laporan_get(id):
    rec = lookup_record(id)
    if not rec:
        return jsonify({'error': 'Data tidak ditemukan'}), 404
    return jsonify(laporan_public(rec))


@app.route('/api/laporan', methods=['POST'])
@login_required
def api_laporan_create():
    payload = request.get_json(silent=True) or {}
    data = parse_payload(payload)
    if not data['kebun'] and not data['kegiatan']:
        return jsonify({'error': 'Isi kebun/lokasi atau kegiatan'}), 400

    conn = get_connection()
    try:
        cur = conn.execute(
            """INSERT INTO laporan
               (kode_laporan, no, regional, perusahaan, kebun, petugas,
                tanggal_kunjungan, kegiatan, draft_masuk, draft_korektor,
                korektor, revisi, cetak, sudah_dikirim, tahun, folder_laporan, catatan)
               VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            (data['kode_laporan'], data['no'], data['regional'],
             data['perusahaan'], data['kebun'], data['petugas'],
             data['tanggal_kunjungan'], data['kegiatan'], data['draft_masuk'],
             data['draft_korektor'], data['korektor'], data['revisi'],
             data['cetak'], data['sudah_dikirim'], data['tahun'],
             data['folder_laporan'], data['catatan'])
        )
        new_id = cur.lastrowid
        conn.commit()
    finally:
        conn.close()

    if not data['kode_laporan'] and new_id:
        _backfill_single(new_id)

    for name in [n.strip() for n in data['korektor'].split(';') if n.strip()]:
        insert_korektor_master(name)

    return jsonify({'status': 'ok', 'id': new_id}), 201


@app.route('/api/laporan/<int:id>', methods=['PUT'])
@login_required
def api_laporan_update(id):
    existing = lookup_record(id)
    if not existing:
        return jsonify({'error': 'Data tidak ditemukan'}), 404

    payload = request.get_json(silent=True) or {}
    data = parse_payload(payload, existing)
    if not data['kebun'] and not data['kegiatan']:
        return jsonify({'error': 'Isi kebun/lokasi atau kegiatan'}), 400

    conn = get_connection()
    try:
        conn.execute(
            """UPDATE laporan SET
               kode_laporan=?, no=?, regional=?, perusahaan=?, kebun=?,
               petugas=?, tanggal_kunjungan=?, kegiatan=?, draft_masuk=?,
               draft_korektor=?, korektor=?, revisi=?, cetak=?, sudah_dikirim=?,
               tahun=?, folder_laporan=?, catatan=?
               WHERE id=?""",
            (data['kode_laporan'], data['no'], data['regional'],
             data['perusahaan'], data['kebun'], data['petugas'],
             data['tanggal_kunjungan'], data['kegiatan'], data['draft_masuk'],
             data['draft_korektor'], data['korektor'], data['revisi'],
             data['cetak'], data['sudah_dikirim'], data['tahun'],
             data['folder_laporan'], data['catatan'], id)
        )
        conn.commit()
    finally:
        conn.close()

    for name in [n.strip() for n in data['korektor'].split(';') if n.strip()]:
        insert_korektor_master(name)

    return jsonify({'status': 'ok'})


@app.route('/api/laporan/<int:id>', methods=['DELETE'])
@login_required
def api_laporan_delete(id):
    existing = lookup_record(id)
    if not existing:
        return jsonify({'error': 'Data tidak ditemukan'}), 404
    conn = get_connection()
    try:
        conn.execute("DELETE FROM laporan WHERE id=?", (id,))
        conn.commit()
    finally:
        conn.close()
    return jsonify({'status': 'ok'})


# ---------------- KOREKTOR ASSIGNMENT API (tahap koreksi) ----------------

@app.route('/api/laporan/<int:id>/korektor_assignment', methods=['GET'])
def api_korektor_assignment_get(id):
    rec = lookup_record(id)
    if not rec:
        return jsonify({'error': 'Data tidak ditemukan'}), 404
    conn = get_connection()
    try:
        rows = conn.execute(
            "SELECT id, urutan, nama, masuk, keluar FROM korektor_assignment WHERE laporan_id=? ORDER BY urutan",
            (id,)).fetchall()
        return jsonify([dict(r) for r in rows])
    finally:
        conn.close()


@app.route('/api/laporan/<int:id>/korektor_assignment', methods=['PUT'])
@login_required
def api_korektor_assignment_put(id):
    rec = lookup_record(id)
    if not rec:
        return jsonify({'error': 'Data tidak ditemukan'}), 404
    payload = request.get_json(silent=True) or {}
    items = payload.get('items')
    if not isinstance(items, list):
        return jsonify({'error': 'Format items harus berupa daftar'}), 400

    conn = get_connection()
    try:
        conn.execute("DELETE FROM korektor_assignment WHERE laporan_id=?", (id,))
        urutan = 0
        for it in items:
            nama = clean_str(it.get('nama'))
            if not nama:
                continue
            urutan += 1
            conn.execute(
                "INSERT INTO korektor_assignment (laporan_id, urutan, nama, masuk, keluar) VALUES (?,?,?,?,?)",
                (id, urutan, nama, clean_str(it.get('masuk')), clean_str(it.get('keluar')))
            )
            insert_korektor_master(nama)
        conn.commit()
    finally:
        conn.close()
    return jsonify({'status': 'ok'})


# ---------------- INDEX LAPORAN API ----------------

@app.route('/api/index_laporan')
def api_index_laporan():
    tahun = request.args.get('tahun', '')
    kategori = request.args.get('kategori', '')
    search = request.args.get('search', '').lower()
    conn = get_connection()
    try:
        sql = "SELECT * FROM index_laporan WHERE 1=1"
        params = []
        if tahun:
            sql += " AND tahun=?"
            params.append(tahun)
        if kategori:
            sql += " AND kategori=?"
            params.append(kategori)
        if search:
            sql += " AND (folder LIKE ? OR sumber LIKE ?)"
            params.append('%' + search + '%')
            params.append('%' + search + '%')
        rows = conn.execute(sql + " ORDER BY folder", params).fetchall()
        return jsonify([dict(r) for r in rows])
    finally:
        conn.close()


# ---------------- EXPORT XLSX ----------------

@app.route('/api/export')
def api_export():
    from io import BytesIO
    import openpyxl
    from openpyxl.styles import PatternFill, Font
    from openpyxl.utils import get_column_letter

    data = query_all_laporan()
    regional = request.args.get('regional', '')
    tahun = request.args.get('tahun', '')
    status = request.args.get('status', '')
    search = request.args.get('search', '').lower()

    filtered = data
    if regional and regional != 'Semua':
        filtered = [d for d in filtered if d['regional'] == regional]
    if tahun and tahun != 'Semua':
        filtered = [d for d in filtered if d['tahun'] == tahun]
    if status and status != 'Semua':
        filtered = [d for d in filtered if d['status'] == status]
    if search:
        filtered = [d for d in filtered if
                    search in str(d['kebun']).lower() or
                    search in str(d['kegiatan']).lower() or
                    search in str(d['korektor']).lower() or
                    search in str(d['perusahaan']).lower() or
                    search in str(d['petugas']).lower() or
                    search in str(d['kode_laporan']).lower() or
                    search in str(d['regional']).lower() or
                    search in str(d['tahun']).lower()]

    cols = [
        ('Kode', 'kode_laporan'), ('No', 'no'), ('Regional', 'regional'),
        ('Perusahaan', 'perusahaan'), ('Kebun/Lokasi', 'kebun'),
        ('Petugas', 'petugas'), ('Tanggal Kunjungan', 'tanggal_kunjungan'),
        ('Kegiatan', 'kegiatan'), ('Draft Masuk', 'draft_masuk'),
        ('Draft Korektor', 'draft_korektor'), ('Korektor', 'korektor'),
        ('Revisi', 'revisi'), ('Cetak', 'cetak'), ('Sudah Dikirim', 'sudah_dikirim'),
        ('Status', 'status'), ('Kategori Kegiatan', 'kegiatan_kategori'),
        ('Tahun', 'tahun'), ('Folder Laporan', 'folder_laporan'), ('Catatan', 'catatan'),
    ]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Data'
    header_fill = PatternFill('solid', fgColor='0d9488')
    header_font = Font(color='FFFFFF', bold=True)
    for c, (label, _key) in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=c, value=label)
        cell.fill = header_fill
        cell.font = header_font
    for r, d in enumerate(filtered, start=2):
        for c, (_label, key) in enumerate(cols, start=1):
            ws.cell(row=r, column=c, value=d.get(key, '') or '')
    for c in range(1, len(cols) + 1):
        ws.column_dimensions[get_column_letter(c)].width = 18

    selesai = sum(1 for d in filtered if d['status'] == 'SELESAI')
    persen = round(selesai / len(filtered) * 100, 1) if filtered else 0
    ws2 = wb.create_sheet('Ringkasan')
    ringkas = [
        ('Total Data', len(filtered)),
        ('Selesai', selesai),
        ('Proses', len(filtered) - selesai),
        ('% Selesai', persen),
    ]
    for i, (k, v) in enumerate(ringkas, start=1):
        ws2.cell(row=i, column=1, value=k)
        ws2.cell(row=i, column=2, value=v)
    ws2.column_dimensions['A'].width = 16

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return (buf.getvalue(), 200,
            {'Content-Type': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
             'Content-Disposition': 'attachment; filename=export_bantek.xlsx'})


# ---------------- KOREKTOR MASTER API ----------------

@app.route('/api/korektor_master', methods=['GET'])
@login_required
def api_korektor_master_list():
    conn = get_connection()
    try:
        rows = conn.execute("SELECT * FROM korektor_master ORDER BY nama").fetchall()
        return jsonify([dict(r) for r in rows])
    finally:
        conn.close()


@app.route('/api/korektor_master', methods=['POST'])
@login_required
def api_korektor_master_create():
    payload = request.get_json(silent=True) or {}
    name = clean_str(payload.get('nama'))
    if not name:
        return jsonify({'error': 'Nama korektor wajib diisi'}), 400
    conn = get_connection()
    try:
        conn.execute("INSERT OR IGNORE INTO korektor_master (nama) VALUES (?)", (name,))
        conn.commit()
    finally:
        conn.close()
    return jsonify({'status': 'ok'}), 201


@app.route('/api/korektor_master/<int:id>', methods=['PUT'])
@login_required
def api_korektor_master_update(id):
    payload = request.get_json(silent=True) or {}
    name = clean_str(payload.get('nama'))
    if not name:
        return jsonify({'error': 'Nama korektor wajib diisi'}), 400
    conn = get_connection()
    try:
        conn.execute("UPDATE korektor_master SET nama=? WHERE id=?", (name, id))
        conn.commit()
    finally:
        conn.close()
    return jsonify({'status': 'ok'})


@app.route('/api/korektor_master/<int:id>', methods=['DELETE'])
@login_required
def api_korektor_master_delete(id):
    conn = get_connection()
    try:
        conn.execute("DELETE FROM korektor_master WHERE id=?", (id,))
        conn.commit()
    finally:
        conn.close()
    return jsonify({'status': 'ok'})


# ---------------- PASSWORD CHANGE ----------------

@app.route('/api/change_password', methods=['POST'])
@login_required
def api_change_password():
    payload = request.get_json(silent=True) or {}
    old = payload.get('old_password', '')
    new = payload.get('new_password', '')
    if not check_admin_password(old):
        return jsonify({'error': 'Password lama salah'}), 400
    if not new or len(new) < 4:
        return jsonify({'error': 'Password baru minimal 4 karakter'}), 400
    hashed = generate_password_hash(new)
    conn = get_connection()
    try:
        conn.execute("INSERT OR REPLACE INTO settings (key, value) VALUES ('admin_password_hash', ?)", (hashed,))
        conn.commit()
    finally:
        conn.close()
    return jsonify({'status': 'ok'})


if __name__ == '__main__':
    if not db_exists():
        print("Database belum ada. Jalankan dulu: python init_db.py")
        raise SystemExit(1)
    ensure_db()
    migrate_db()
    migrate_and_backfill()
    app.run(debug=True, host='0.0.0.0', port=5000)
