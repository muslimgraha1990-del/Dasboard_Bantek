import os
import sys
import re
import shutil
import datetime

import openpyxl

from database import get_connection, init_db, migrate_db, DB_PATH

EXCEL_PATH = os.path.join(
    os.path.dirname(os.path.abspath(__file__)),
    '..',
    'Monitoring Bantuan Teknis - Ringkasan_KOREKTOR_3.xlsx'
)

DATA_SHEETS = {
    'Data R1': 'Regional 1',
    'Data R2': 'Regional 2',
    'Data R3': 'Regional 3',
    'Data R4': 'Regional 4',
    'Data R5': 'Regional 5',
    'Data R6': 'Regional 6',
    'Data R7': 'Regional 7',
    'Data R4P': 'Regional 4 Palmco',
    'Data SW': 'PT Swasta / PPKS',
}

INDEX_SHEET = 'INDEX LAPORAN'

REGIONAL_NO = {
    'Regional 1': '1', 'Regional 2': '2', 'Regional 3': '3', 'Regional 4': '4',
    'Regional 5': '5', 'Regional 6': '6', 'Regional 7': '7',
    'Regional 4 Palmco': '4P', 'PT Swasta / PPKS': 'SW',
}

MAX_KOREKTOR_SLOT = 8


def norm_date(v):
    if v is None:
        return ''
    if isinstance(v, datetime.datetime):
        return v.strftime('%d-%m-%Y')
    if isinstance(v, datetime.date):
        return v.strftime('%d-%m-%Y')
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        try:
            from openpyxl.utils.datetime import from_excel
            return from_excel(v).strftime('%d-%m-%Y')
        except Exception:
            pass
    s = str(v).strip()
    if not s:
        return ''
    lowered = s.lower()
    if ' 00:00:00' in lowered or lowered.startswith('1900-01-00') or lowered.startswith('1899-'):
        return s
    if len(s) >= 8 and s[4] == '-' and s[7] == '-':
        try:
            return datetime.datetime.strptime(s[:10], '%Y-%m-%d').strftime('%d-%m-%Y')
        except ValueError:
            pass
    return s


def clean(v):
    return str(v).strip() if v is not None else ''


def read_korektor_names(ws):
    names = []
    for i in range(MAX_KOREKTOR_SLOT):
        col = 10 + i * 2
        v = ws.cell(row=2, column=col).value
        if v:
            names.append(clean(v))
    return names


def extract_link(cell):
    if cell.hyperlink is not None and cell.hyperlink.target:
        return clean(cell.hyperlink.target)
    v = cell.value
    if isinstance(v, str) and v.strip().upper().startswith('=HYPERLINK'):
        m = re.match(r'^=HYPERLINK\(\s*"([^"]+)"', v.strip(), re.IGNORECASE)
        if m:
            return m.group(1)
    return clean(v)


def backup_db():
    if not os.path.exists(DB_PATH):
        return None
    ts = datetime.datetime.now().strftime('%Y%m%d-%H%M%S')
    dest = DB_PATH + '.backup-' + ts
    shutil.copy2(DB_PATH, dest)
    return dest


def load_from_excel(path):
    if not os.path.exists(path):
        print("ERROR: File Excel tidak ditemukan: %s" % path)
        sys.exit(1)

    wb = openpyxl.load_workbook(path, data_only=False)
    korektor_master = set()
    assignments = []
    records = []

    for sheet_name, regional_name in DATA_SHEETS.items():
        if sheet_name not in wb.sheetnames:
            print("  Warning: sheet %s tidak ada, dilewati" % sheet_name)
            continue
        ws = wb[sheet_name]
        names = read_korektor_names(ws)
        for n in names:
            korektor_master.add(n)

        for row_idx in range(4, ws.max_row + 1):
            no = ws.cell(row=row_idx, column=1).value
            kebun = ws.cell(row=row_idx, column=3).value
            if no is None or not kebun or clean(kebun) == '':
                continue

            tahun = ws.cell(row=row_idx, column=30).value
            if isinstance(tahun, (int, float)):
                tahun_str = str(int(tahun))
            else:
                tahun_str = clean(tahun)

            row_assignments = []
            for i, name in enumerate(names):
                col_masuk = 10 + i * 2
                masuk = norm_date(ws.cell(row=row_idx, column=col_masuk).value)
                if not masuk:
                    continue
                keluar = norm_date(ws.cell(row=row_idx, column=col_masuk + 1).value)
                row_assignments.append({
                    'urutan': i + 1, 'nama': name, 'masuk': masuk, 'keluar': keluar,
                })
                korektor_master.add(name)

            korektor_display = clean(ws.cell(row=row_idx, column=9).value)
            if not korektor_display and row_assignments:
                korektor_display = '; '.join(a['nama'] for a in row_assignments)

            folder = extract_link(ws.cell(row=row_idx, column=31))

            record = {
                'no': clean(no),
                'regional': regional_name,
                'perusahaan': clean(ws.cell(row=row_idx, column=2).value),
                'kebun': clean(kebun),
                'petugas': clean(ws.cell(row=row_idx, column=4).value),
                'tanggal_kunjungan': clean(ws.cell(row=row_idx, column=5).value),
                'kegiatan': clean(ws.cell(row=row_idx, column=6).value),
                'draft_masuk': norm_date(ws.cell(row=row_idx, column=7).value),
                'draft_korektor': norm_date(ws.cell(row=row_idx, column=8).value),
                'korektor': korektor_display,
                'revisi': norm_date(ws.cell(row=row_idx, column=26).value),
                'cetak': norm_date(ws.cell(row=row_idx, column=27).value),
                'sudah_dikirim': clean(ws.cell(row=row_idx, column=28).value),
                'tahun': tahun_str,
                'folder_laporan': folder,
                'catatan': clean(ws.cell(row=row_idx, column=32).value),
            }
            records.append(record)
            assignments.append({'regional': regional_name, 'no': record['no'], 'items': row_assignments})

    index_records = []
    if INDEX_SHEET in wb.sheetnames:
        ws = wb[INDEX_SHEET]
        start_row = 2
        if clean(ws.cell(row=1, column=1).value).lower() == 'folder':
            start_row = 2
        else:
            start_row = 1
        for row_idx in range(start_row, ws.max_row + 1):
            folder = clean(ws.cell(row=row_idx, column=1).value)
            if not folder:
                continue
            link = extract_link(ws.cell(row=row_idx, column=6))
            index_records.append({
                'folder': folder,
                'tahun': clean(ws.cell(row=row_idx, column=2).value),
                'kategori': clean(ws.cell(row=row_idx, column=3).value),
                'sumber': clean(ws.cell(row=row_idx, column=4).value),
                'jumlah_file': clean(ws.cell(row=row_idx, column=5).value),
                'link': link,
            })

    wb.close()
    return records, assignments, index_records, sorted(korektor_master)


def build_lookup(assignments):
    lookup = {}
    for a in assignments:
        lookup.setdefault((a['regional'], a['no']), []).extend(a['items'])
    return lookup


def excel_summary(path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    out = {}
    for sheet_name, regional_name in DATA_SHEETS.items():
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        total = 0
        selesai = 0
        for row in ws.iter_rows(min_row=4, min_col=1, max_col=28, values_only=True):
            no, kebun, ab = row[0], row[2], row[27]
            if no is not None and kebun and clean(kebun):
                total += 1
                if ab is not None and str(ab).strip():
                    selesai += 1
        out[regional_name] = {'total': total, 'selesai': selesai}

    index_count = 0
    if INDEX_SHEET in wb.sheetnames:
        ws = wb[INDEX_SHEET]
        start_row = 2 if clean(ws.cell(row=1, column=1).value).lower() == 'folder' else 1
        for row_idx in range(start_row, ws.max_row + 1):
            if clean(ws.cell(row=row_idx, column=1).value):
                index_count += 1
    wb.close()
    return out, index_count


def db_summary():
    conn = get_connection()
    try:
        rows = conn.execute(
            "SELECT regional, COUNT(*) c, "
            "SUM(CASE WHEN sudah_dikirim IS NOT NULL AND sudah_dikirim!='' THEN 1 ELSE 0 END) s "
            "FROM laporan GROUP BY regional").fetchall()
        out = {r['regional']: {'total': r['c'], 'selesai': r['s'] or 0} for r in rows}
        idx = conn.execute("SELECT COUNT(*) c FROM index_laporan").fetchone()['c']
        return out, idx
    finally:
        conn.close()


def compare_report(path):
    ecc, eidx = excel_summary(path)
    dcc, didx = db_summary()
    detail = {}
    for reg in DATA_SHEETS.values():
        e = ecc.get(reg, {'total': 0, 'selesai': 0})
        d = dcc.get(reg, {'total': 0, 'selesai': 0})
        detail[reg] = {
            'excel_total': e['total'], 'db_total': d['total'],
            'excel_selesai': e['selesai'], 'db_selesai': d['selesai'],
            'match': e['total'] == d['total'] and e['selesai'] == d['selesai'],
        }
    match = all(v['match'] for v in detail.values()) and \
        sum(v['total'] for v in ecc.values()) == sum(v['total'] for v in dcc.values()) and \
        eidx == didx
    return {
        'file': os.path.basename(path),
        'file_mtime': datetime.datetime.fromtimestamp(os.path.getmtime(path)).strftime('%Y-%m-%d %H:%M:%S'),
        'excel_total': sum(v['total'] for v in ecc.values()),
        'db_total': sum(v['total'] for v in dcc.values()),
        'excel_index': eidx,
        'db_index': didx,
        'match': match,
        'detail': detail,
    }


def rebuild_db(records, assignments, index_records, korektor_master):
    conn = get_connection()
    try:
        conn.execute("DELETE FROM korektor_assignment")
        conn.execute("DELETE FROM index_laporan")
        conn.execute("DELETE FROM laporan")
        conn.execute("DELETE FROM sqlite_sequence WHERE name IN ('laporan','korektor_assignment','index_laporan')")

        seqs = {}
        inserted = 0
        for r in records:
            cur = conn.execute(
                """INSERT INTO laporan
                   (kode_laporan, no, regional, perusahaan, kebun, petugas,
                    tanggal_kunjungan, kegiatan, draft_masuk, draft_korektor,
                    korektor, revisi, cetak, sudah_dikirim, tahun, folder_laporan, catatan)
                   VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                (r['kode_laporan'], r['no'], r['regional'], r['perusahaan'], r['kebun'],
                 r['petugas'], r['tanggal_kunjungan'], r['kegiatan'], r['draft_masuk'],
                 r['draft_korektor'], r['korektor'], r['revisi'], r['cetak'],
                 r['sudah_dikirim'], r['tahun'], r['folder_laporan'], r['catatan'])
            )
            new_id = cur.lastrowid

            no_reg = REGIONAL_NO.get(r['regional'], '0')
            key = (r['regional'], r['tahun'])
            seqs[key] = seqs.get(key, 0) + 1
            tahun = r['tahun'] or '0000'
            kode = 'R%s-%s-%03d' % (no_reg, tahun, seqs[key])
            conn.execute("UPDATE laporan SET kode_laporan=? WHERE id=?", (kode, new_id))

            for a in assignments.get((r['regional'], r['no']), []):
                conn.execute(
                    """INSERT INTO korektor_assignment
                       (laporan_id, urutan, nama, masuk, keluar)
                       VALUES (?,?,?,?,?)""",
                    (new_id, a['urutan'], a['nama'], a['masuk'], a['keluar'])
                )
            inserted += 1

        keep = set(korektor_master)
        for r in conn.execute("SELECT DISTINCT nama FROM korektor_assignment"):
            keep.add(r['nama'])
        for r in conn.execute("SELECT DISTINCT korektor FROM laporan WHERE korektor IS NOT NULL AND korektor != ''"):
            for name in [n.strip() for n in str(r['korektor']).split(';') if n.strip()]:
                keep.add(name)
        conn.execute("DELETE FROM korektor_master")
        for name in sorted(keep):
            conn.execute("INSERT OR IGNORE INTO korektor_master (nama) VALUES (?)", (name,))

        for it in index_records:
            conn.execute(
                "INSERT INTO index_laporan (folder, tahun, kategori, sumber, jumlah_file, link) VALUES (?,?,?,?,?,?)",
                (it['folder'], it['tahun'], it['kategori'], it['sumber'], it['jumlah_file'], it['link'])
            )

        conn.commit()
        return inserted
    finally:
        conn.close()


def import_from_excel(path):
    backup = backup_db()
    init_db()
    migrate_db()
    records, assignments, index_records, korektor_master = load_from_excel(path)
    lookup = build_lookup(assignments)
    for r in records:
        r['kode_laporan'] = ''
    inserted = rebuild_db(records, lookup, index_records, korektor_master)
    return {
        'inserted': inserted,
        'korektor': len(korektor_master),
        'index': len(index_records),
        'backup': os.path.basename(backup) if backup else None,
    }


def main():
    path = EXCEL_PATH
    if len(sys.argv) > 1 and sys.argv[1] != '--sync':
        path = sys.argv[1]

    print("=" * 60)
    print("Import data dari %s" % os.path.basename(path))
    print("=" * 60)

    before = compare_report(path) if os.path.exists(DB_PATH) else None

    result = import_from_excel(path)

    after = compare_report(path)
    print("  Backup DB   : %s" % (result['backup'] or '(tidak ada)'))
    print("  Laporan     : %d data" % result['inserted'])
    print("  Korektor    : %d nama" % result['korektor'])
    print("  Index folder: %d baris" % result['index'])
    print("=" * 60)
    if before and not before.get('match'):
        print("Peringatan: sebelum sync, data DB berbeda dari Excel.")
    if after.get('match'):
        print("Status     : DATABASE SINKRON DENGAN EXCEL [OK]")
    else:
        print("Status     : MASIH ADA PERBEDAAN - periksa laporan per regional")
    print("Selesai. Jalankan: python app.py")


if __name__ == '__main__':
    main()